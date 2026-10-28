import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\pc\Downloads\Airborne Feedback Sheet.xlsx')
ws = wb['Table 1']

print("Dimensions:", ws.dimensions)

for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(vals):
        task_id = str(vals[0] or '')
        if 'CRM' in task_id or 'SMS' in task_id or 'CMS' in task_id or r == 1:
            print(f"Row {r}: {vals[:9]}")
