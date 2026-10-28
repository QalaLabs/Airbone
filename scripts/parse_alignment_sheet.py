import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\pc\Downloads\Airborne Feedback & Codebase Alignment.xlsx', data_only=True)
ws = wb['Feedback Tracker']

headers = [cell.value for cell in ws[1]]
print("Headers in Feedback Tracker:", headers)

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if any(row):
        row_dict = dict(zip(headers, row))
        rows.append(row_dict)

print(f"Total rows in Feedback Tracker: {len(rows)}")

crm_rows = [r for r in rows if str(r.get('Task ID', '')).startswith('CRM')]
print(f"Total CRM rows: {len(crm_rows)}")
for r in crm_rows:
    print(f"[{r.get('Task ID')}] [{r.get('Priority')}] [{r.get('Status')}] {r.get('Task Title')}")

with open('scripts/extracted_alignment_sheet.json', 'w', encoding='utf-8') as f:
    json.dump(rows, f, indent=2, default=str)
